import json
import urllib.request
import os
import glob
import openpyxl

# 1. Load env
env_files = [".env.production.local", ".env.local", ".env.production", ".env"]
env = {}
for ef in env_files:
    p = os.path.join("d:/orange-msk-main", ef)
    if os.path.exists(p):
        with open(p, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if line and not line.startswith("#") and "=" in line:
                    k, v = line.split("=", 1)
                    if k.strip() not in env:
                        env[k.strip()] = v.strip().strip("'\"")

url = env.get("SUPABASE_URL") or env.get("NEXT_PUBLIC_SUPABASE_URL")
key = env.get("SUPABASE_SERVICE_ROLE_KEY") or env.get("NEXT_PUBLIC_SUPABASE_ANON_KEY")
base_url = "https://orangemsk.ru"

# 2. Fetch products
req = urllib.request.Request(
    f"{url}/rest/v1/products?select=*&is_visible=eq.true&order=id.asc",
    headers={"apikey": key, "Authorization": f"Bearer {key}"}
)

with urllib.request.urlopen(req) as resp:
    products = json.loads(resp.read().decode("utf-8"))

print(f"Loaded {len(products)} products from Supabase.")

# 3. Classification helper
def classify(p):
    name = (p.get("name") or "").lower()
    cat = (p.get("category") or "").lower()
    brand = (p.get("brand") or "").strip()
    
    market_cat = "электроника / телефоны / смартфоны"
    tn_ved = "8517130000"
    weight = 0.45
    length = 18.0
    width = 10.0
    height = 4.0
    country = "Китай"
    life = "3 года"
    
    if "lego" in name or cat == "lego" or brand.lower() == "lego":
        market_cat = "детские товары / игрушки и игры / конструкторы / конструкторы и радиоконструкторы / конструкторы детские"
        tn_ved = "9503003500"
        weight = 1.2
        length = 38.0
        width = 26.0
        height = 9.0
        country = "Дания"
        brand = "LEGO"
        life = "5 лет"
    elif "dyson" in name or cat == "dyson" or brand.lower() == "dyson":
        brand = "Dyson"
        country = "Малайзия"
        if "пылесос" in name:
            market_cat = "бытовая техника / техника для дома / техника для уборки / вертикальные пылесосы"
            tn_ved = "8508110000"
            weight = 4.5
            length = 90.0
            width = 30.0
            height = 25.0
        elif "airstrait" in name or "выпрямитель" in name or "стайлер" in name or "corrale" in name:
            market_cat = "бытовая техника / техника для красоты / щипцы, плойки и выпрямители / выпрямители для укладки волос"
            tn_ved = "8516320000"
            weight = 1.8
            length = 40.0
            width = 20.0
            height = 12.0
        else:
            market_cat = "бытовая техника / техника для красоты / фены и фен-щётки / фены для сушки волос"
            tn_ved = "8516310000"
            weight = 1.9
            length = 38.0
            width = 24.0
            height = 14.0
    elif "airpods" in name or "наушники" in name:
        market_cat = "электроника / портативная техника / наушники"
        tn_ved = "8518300000"
        brand = "Apple"
        if "max" in name:
            weight = 0.85
            length = 24.0
            width = 24.0
            height = 8.0
        else:
            weight = 0.25
            length = 10.0
            width = 10.0
            height = 5.0
    elif "macbook" in name or "ноутбук" in name:
        market_cat = "компьютерная техника / компьютеры / ноутбуки персональные"
        tn_ved = "8471300000"
        brand = "Apple"
        weight = 2.2
        length = 36.0
        width = 27.0
        height = 6.0
    elif "ipad" in name or "планшет" in name or "galaxy tab" in name:
        market_cat = "компьютерная техника / компьютеры / планшетные компьютеры"
        tn_ved = "8471300000"
        if "galaxy" in name or "samsung" in name:
            brand = "Samsung"
            country = "Вьетнам"
        else:
            brand = "Apple"
        weight = 0.85
        length = 28.0
        width = 21.0
        height = 4.5
    elif "watch" in name or "часы" in name:
        market_cat = "электроника / телефоны / умные часы и браслеты / умные часы"
        tn_ved = "8517620003"
        brand = "Apple"
        weight = 0.35
        length = 25.0
        width = 8.0
        height = 4.0
    elif "колонки" in name or "колонка" in name or "speaker" in name or "homepod" in name:
        market_cat = "электроника / портативная техника / беспроводные колонки"
        tn_ved = "8518210000"
        weight = 1.5
        length = 20.0
        width = 20.0
        height = 20.0
    elif "playstation" in name or "xbox" in name or "приставка" in name or "nintendo" in name:
        market_cat = "электроника / игровые приставки и аксессуары / приставки игровые"
        tn_ved = "9504500001"
        weight = 4.5
        length = 45.0
        width = 35.0
        height = 15.0
    elif "аккумулятор" in name or "power bank" in name or "battery" in name or "magsafe battery" in name:
        market_cat = "электроника / телефоны / аксессуары для телефонов / портативные аккумуляторы"
        tn_ved = "8507600000"
        weight = 0.35
        length = 15.0
        width = 10.0
        height = 3.5
    elif "чехол" in name or "case" in name or "кабель" in name or "cable" in name or "адаптер" in name or "adapter" in name or cat == "accessories":
        market_cat = "компьютерная техника / аксессуары / кабели, разъемы, переходники для компьютеров и электроники / кабели для мобильных устройств"
        tn_ved = "8544429007"
        weight = 0.15
        length = 15.0
        width = 10.0
        height = 3.0
    else: # smartphone
        market_cat = "электроника / телефоны / смартфоны"
        tn_ved = "8517130000"
        weight = 0.45
        length = 18.0
        width = 10.0
        height = 4.0
        if "samsung" in name or brand.lower() == "samsung":
            country = "Вьетнам"
            brand = "Samsung"
        elif "apple" in name or "iphone" in name or brand.lower() == "apple":
            brand = "Apple"
            
    return {
        "market_cat": market_cat,
        "brand": brand,
        "tn_ved": tn_ved,
        "weight": weight,
        "length": length,
        "width": width,
        "height": height,
        "country": country,
        "life": life,
    }

# 4. Use pristine template
template_path = r"C:\Users\Artem\.gemini\antigravity\brain\cd9308ac-3ace-4887-92a1-aa1bb1bccb23\.user_uploaded\media_1791407881203.xlsx"
if not os.path.exists(template_path):
    # fallback
    template_path = r"C:\Users\Artem\.gemini\antigravity\brain\cd9308ac-3ace-4887-92a1-aa1bb1bccb23\.user_uploaded\media_1791408398598.xlsx"

print(f"Loading template from: {template_path}")
wb = openpyxl.load_workbook(template_path)

# Find the "Список товаров" worksheet: it's worksheet index 2 (has 48 columns)
ws = None
for s in wb.worksheets:
    if s.cell(2, 4).value and "SKU" in str(s.cell(2, 4).value):
        ws = s
        break

if ws is None:
    ws = wb.worksheets[2]

print(f"Selected worksheet for products: {ws.title} (cols: {ws.max_column}, rows: {ws.max_row})")

# Delete existing data rows starting from row 4
if ws.max_row >= 4:
    ws.delete_rows(4, ws.max_row - 3)

# Filter out test products or products without price / dummy
valid_products = []
for p in products:
    name = (p.get("name") or "").strip()
    price = int(p.get("price") or 0)
    if "тест" in name.lower() or "проверки оплаты" in name.lower():
        print(f"Skipping test product: {name}")
        continue
    if price <= 100:
        print(f"Skipping product with price <= 100: {name}")
        continue
    valid_products.append(p)

print(f"Exporting {len(valid_products)} products to 'Список товаров'...")

start_row = 4
for idx, p in enumerate(valid_products):
    curr_row = start_row + idx
    meta = classify(p)
    
    sku = p.get("sku") or p.get("code") or str(p["id"])
    name = (p.get("name") or "").strip()
    
    # Process images
    raw_images = p.get("images") or []
    if isinstance(raw_images, str):
        try:
            raw_images = json.loads(raw_images)
        except:
            raw_images = [raw_images]
    
    clean_images = []
    for img in raw_images:
        if isinstance(img, str) and img.strip():
            img_clean = img.strip()
            if img_clean.startswith("/"):
                clean_images.append(f"{base_url}{img_clean}")
            elif img_clean.startswith("http"):
                clean_images.append(img_clean)
    
    images_str = ", ".join(clean_images[:30])
    description = (p.get("description") or name).strip()
    if len(description) > 5900:
        description = description[:5900]
    
    barcode = p.get("code") or sku
    price = int(p.get("price") or 0)
    old_price = int(p["old_price"]) if p.get("old_price") and p["old_price"] > price else ""
    vol = round(meta["length"] * meta["width"] * meta["height"] / 1000.0, 3)
    
    # Fill row
    ws.cell(curr_row, 1, "") # Критичные ошибки
    ws.cell(curr_row, 2, "") # Некритичные ошибки
    ws.cell(curr_row, 3, "") # Качество карточки
    ws.cell(curr_row, 4, str(sku)) # Ваш SKU *
    ws.cell(curr_row, 5, name) # Название товара *
    ws.cell(curr_row, 6, images_str) # Ссылка на изображение *
    ws.cell(curr_row, 7, description) # Описание товара *
    ws.cell(curr_row, 8, meta["market_cat"]) # Категория на Маркете *
    ws.cell(curr_row, 9, meta["brand"]) # Бренд *
    ws.cell(curr_row, 10, str(barcode)) # Штрихкод *
    ws.cell(curr_row, 11, p.get("category") or "") # Теги
    ws.cell(curr_row, 12, "") # Ссылка на видео
    ws.cell(curr_row, 13, "") # Инструкции
    ws.cell(curr_row, 14, meta["country"]) # Страна производства
    ws.cell(curr_row, 15, str(sku)) # Артикул производителя
    ws.cell(curr_row, 16, meta["weight"]) # Вес, кг *
    ws.cell(curr_row, 17, meta["length"]) # Длина, см *
    ws.cell(curr_row, 18, meta["width"]) # Ширина, см *
    ws.cell(curr_row, 19, meta["height"]) # Высота, см *
    ws.cell(curr_row, 20, "") # Товар доставляется в нескольких упаковках
    ws.cell(curr_row, 21, vol) # Объём, л
    ws.cell(curr_row, 22, price) # Цена *
    ws.cell(curr_row, 23, old_price) # Зачёркнутая цена
    ws.cell(curr_row, 24, "") # Себестоимость
    ws.cell(curr_row, 25, "") # Дополнительные расходы
    ws.cell(curr_row, 26, "") # Срок годности
    ws.cell(curr_row, 27, "") # Комментарий к сроку годности
    ws.cell(curr_row, 28, meta["life"]) # Срок службы
    ws.cell(curr_row, 29, "") # Комментарий к сроку службы
    ws.cell(curr_row, 30, "1 год") # Гарантийный срок
    ws.cell(curr_row, 31, "") # Комментарий к гарантийному сроку
    ws.cell(curr_row, 32, "") # Маркировка
    ws.cell(curr_row, 33, meta["tn_ved"]) # ТН ВЭД *
    ws.cell(curr_row, 34, "") # ОКПД 2
    ws.cell(curr_row, 35, "") # Номер документа на товар
    ws.cell(curr_row, 36, "") # Тип уценки
    ws.cell(curr_row, 37, "") # Внешний вид товара
    ws.cell(curr_row, 38, "") # Описание состояния товара
    ws.cell(curr_row, 39, "") # Особый тип товара
    ws.cell(curr_row, 40, "") # С какого возраста пользоваться
    ws.cell(curr_row, 41, "Нет") # Товар для взрослых
    ws.cell(curr_row, 42, "Нет") # Цифровой товар
    ws.cell(curr_row, 43, "") # Характеристики товара
    ws.cell(curr_row, 44, "Нет") # В архиве
    ws.cell(curr_row, 45, "") # Артикул товара (SKU) -> market-sku на Маркете! ДОЛЖЕН БЫТЬ ПУСТЫМ!
    ws.cell(curr_row, 46, "") # Артикул Маркета
    ws.cell(curr_row, 47, "") # Категория на Маркете
    ws.cell(curr_row, 48, "") # Дата дополнения карточки

out1 = "C:/Users/Artem/Downloads/Каталог_OrangeMSK_для_Яндекс_Маркета.xlsx"
out2 = "C:/Users/Artem/Downloads/orange_msk_market.xlsx"
out3 = "d:/orange-msk-main/Каталог_OrangeMSK_для_Яндекс_Маркета.xlsx"
out4 = "d:/orange-msk-main/public/yandex_market_catalog_ready.xlsx"

wb.save(out1)
wb.save(out2)
wb.save(out3)
wb.save(out4)

print("SUCCESS: Catalog generated cleanly into all target locations!")
